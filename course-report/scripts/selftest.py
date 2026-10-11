#!/usr/bin/env python3
"""自检：用 tests/sample.md 跑一遍完整流程，逐项检查结构、版面和（有预览 PDF 时）页面。把整段输出贴回即可。

用法：python scripts/selftest.py                  # 页面检查做不了时记为 SKIP，不算失败
      python scripts/selftest.py --require-pages  # SKIP 也算失败
退出码：0 = 没有 FAIL（可能有 SKIP，见最后一行）；1 = 有 FAIL，或 --require-pages 时有 SKIP
"""

import re
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

SKILL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL / "scripts"))
from toolchain import pdf_text_pages  # noqa: E402
results = []  # "PASS" / "FAIL" / "SKIP"


def report(status, name, detail=""):
    results.append(status)
    print("%s  %s%s" % (status, name, ("  —  " + str(detail)) if detail else ""))


def check(name, ok, detail=""):
    report("PASS" if ok else "FAIL", name, detail)


def skip(name, reason):
    report("SKIP", name, reason)


def text_of(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def last_line(page_text):
    lines = [l.strip() for l in page_text.splitlines() if l.strip()]
    return lines[-1] if lines else ""


def main():
    out = Path(tempfile.mkdtemp(prefix="cr_selftest_"))
    r = subprocess.run([sys.executable, str(SKILL / "scripts" / "build_report.py"),
                        str(SKILL / "tests" / "sample.md"), "--out", str(out)],
                       capture_output=True, text=True, encoding="utf-8")
    print(r.stdout)
    if r.returncode != 0:
        print(r.stderr[-1500:])
    docx_path = out / "sample.docx"
    check("构建成功（退出码 0 且生成了 docx）", r.returncode == 0 and docx_path.exists(),
          "退出码 %s，docx %s" % (r.returncode, "存在" if docx_path.exists() else "不存在"))
    if not docx_path.exists():
        summary(out)

    doc = Document(docx_path)
    secs = doc.sections
    check("分成 3 节（封面/目录摘要/正文）", len(secs) == 3, len(secs))

    body = list(doc.element.body.iterchildren())
    first_sect_idx = next((i for i, e in enumerate(body)
                           if e.tag == qn("w:p") and e.find(qn("w:pPr")) is not None
                           and e.find(qn("w:pPr")).find(qn("w:sectPr")) is not None), 0)
    cover_text = "".join(text_of(e) for e in body[:first_sect_idx + 1])
    check("封面含姓名", "张三" in cover_text)
    check("封面含学号", "2024000000000" in cover_text)
    check("封面含写作日期（自动）", re.search(r"\d{4}年\d{1,2}月\d{1,2}日", cover_text), re.findall(r"\d{4}年\d{1,2}月\d{1,2}日", cover_text))

    def has_page_field(section):
        return any("PAGE" in (t.text or "") for t in section.footer._element.iter(qn("w:instrText")))

    check("封面无页码、其余两节有页码", not has_page_field(secs[0]) and has_page_field(secs[1]) and has_page_field(secs[2]))
    fmt = []
    for s in secs:
        pg = s._sectPr.find(qn("w:pgNumType"))
        fmt.append((pg.get(qn("w:fmt")), pg.get(qn("w:start"))) if pg is not None else None)
    check("目录摘要节为罗马数字、正文节从 1 开始", fmt[1] == ("lowerRoman", "1") and fmt[2] == ("decimal", "1"), fmt)
    cols = [(s._sectPr.find(qn("w:cols")).get(qn("w:num")) if s._sectPr.find(qn("w:cols")) is not None else "1") for s in secs]
    check("所有分节单栏", all(c == "1" for c in cols), cols)

    titles = [p for p in doc.paragraphs if p.style.name == "Abstract Title"]
    check("摘要标题套摘要样式（中英各一）", len(titles) == 2, [t.text for t in titles])
    abs_shd = 0
    for p in doc.paragraphs:
        if p.style.name in ("Abstract Title", "Abstract Text", "Keywords"):
            abs_shd += sum(1 for _ in p._p.iter(qn("w:shd")))
    check("摘要没有底纹", abs_shd == 0, abs_shd)
    check("摘要没有外框表格", not any(len(t.rows) == 1 and len(t.columns) == 1 for t in doc.tables))

    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    check("一级标题编号为“一、”", h1 and h1[0].startswith("一、"), h1[:1])
    check("二级标题编号为“（一）”", h2 and h2[0].startswith("（一）"), h2[:1])
    check("参考文献不编号", "参考文献" in h1, h1[-1:])

    toc = [p for p in doc.paragraphs if p.style.name in ("toc 1", "toc 2", "toc 3")]
    heads = [p for p in doc.paragraphs if re.match(r"^Heading [123]$", p.style.name)]
    check("目录条目数与 1–3 级标题数一致", len(toc) == len(heads), "%d vs %d" % (len(toc), len(heads)))
    instr = [t.text for t in doc.element.body.iter(qn("w:instrText"))]
    check("目录是 TOC 域", any("TOC" in (t or "") for t in instr))

    def border_ok(tbl):
        b = tbl._tbl.tblPr.find(qn("w:tblBorders"))
        return b is not None and b.find(qn("w:top")).get(qn("w:val")) == "single" and b.find(qn("w:insideV")).get(qn("w:val")) == "nil"
    check("表格为三线表", doc.tables and all(border_ok(t) for t in doc.tables if len(t.columns) > 2))

    layout_checks(doc, body[:first_sect_idx + 1])

    pdf = out / "preview" / "sample_preview.pdf"
    if pdf.exists():
        pages = pdf_text_pages(pdf) or [""]
        leader = re.compile(r"(\.\s*){4,}")
        check("封面页没有页码", not re.fullmatch(r"[ivx\d]+", last_line(pages[0])), last_line(pages[0]))
        check("目录页页码为 i", last_line(pages[1]) == "i", last_line(pages[1]))
        first_body = next((i for i, p in enumerate(pages) if any(l.strip().startswith("一、") and not leader.search(l) for l in p.splitlines())), None)
        check("正文首页页码为 1", first_body is not None and last_line(pages[first_body]) == "1",
              last_line(pages[first_body]) if first_body is not None else "找不到正文首页")
        nums = re.findall(r"\.{4,}\s*(\d+)\s*$", pages[1], re.M)
        check("目录页码回填为数字", len(nums) == len(toc), "%d 条" % len(nums))
    else:
        for name in ("封面页没有页码", "目录页页码为 i", "正文首页页码为 1", "目录页码回填为数字"):
            skip(name, "没有预览 PDF（无可用渲染器）")

    summary(out)


def summary(out):
    n = {k: results.count(k) for k in ("PASS", "FAIL", "SKIP")}
    if n["FAIL"]:
        verdict = "有失败项"
    elif n["SKIP"]:
        verdict = "未完成：%d 项跳过，不能视为全部通过" % n["SKIP"]
    else:
        verdict = "全部通过"
    print("\nSELFTEST RESULT: %d 通过 / %d 失败 / %d 跳过 —— %s  输出目录：%s"
          % (n["PASS"], n["FAIL"], n["SKIP"], verdict, out))
    strict = "--require-pages" in sys.argv
    sys.exit(1 if n["FAIL"] or (strict and n["SKIP"]) else 0)


# ---------------------------------------------------------------- 版面检查
# Word 的缩进按属性组逐层继承：段落 > 编号级别 > 段落样式链 > 文档默认。
# “字符”组（firstLineChars/hangingChars）非 0 时优先于长度组（firstLine/hanging）。
GROUPS = {"first_chars": ("firstLineChars", "hangingChars"), "first_tw": ("firstLine", "hanging"),
          "left_chars": ("leftChars", "startChars"), "left_tw": ("left", "start")}


def numbering_ind(doc, num_pr):
    num_id = num_pr.find(qn("w:numId"))
    ilvl = num_pr.find(qn("w:ilvl"))
    level = ilvl.get(qn("w:val")) if ilvl is not None else "0"
    if num_id is None:
        return None
    numbering = doc.part.numbering_part.element
    abs_id = next((n.find(qn("w:abstractNumId")).get(qn("w:val")) for n in numbering.findall(qn("w:num"))
                   if n.get(qn("w:numId")) == num_id.get(qn("w:val"))), None)
    for absnum in numbering.findall(qn("w:abstractNum")):
        if absnum.get(qn("w:abstractNumId")) == abs_id:
            for lvl in absnum.findall(qn("w:lvl")):
                if lvl.get(qn("w:ilvl")) == level:
                    return lvl.find(".//" + qn("w:ind"))
    return None


def ind_layers(doc, p):
    layers = []
    ppr = p._p.pPr
    if ppr is not None:
        layers.append(ppr.find(qn("w:ind")))
        num_pr = ppr.find(qn("w:numPr"))
        if num_pr is not None:
            layers.append(numbering_ind(doc, num_pr))
    style = p.style
    while style is not None:
        spr = style.element.pPr
        layers.append(spr.find(qn("w:ind")) if spr is not None else None)
        style = style.base_style
    dflt = doc.styles.element.find(qn("w:docDefaults"))
    if dflt is not None:
        layers.append(dflt.find(".//" + qn("w:ind")))
    return [layer for layer in layers if layer is not None]


def group_value(layers, group):
    """取最近一层定义了该组任一属性的值；hanging 记为负数。"""
    pos, neg = GROUPS[group]
    for ind in layers:
        if ind.get(qn("w:" + pos)) is not None:
            return int(ind.get(qn("w:" + pos)))
        if ind.get(qn("w:" + neg)) is not None:
            return -int(ind.get(qn("w:" + neg)))
    return 0


def effective_chars(doc, p, which):
    """有效缩进（字符；长度单位按 12pt 正文折算）。which = first | left"""
    layers = ind_layers(doc, p)
    chars = group_value(layers, which + "_chars")
    return chars / 100 if chars else group_value(layers, which + "_tw") / 240


def text_len_cm(text, size_pt):
    ems = sum(1.0 if unicodedata.east_asian_width(ch) in "WF" else 0.55 for ch in text)
    return ems * size_pt / 72 * 2.54


def layout_checks(doc, cover_els):
    sec = doc.sections[-1]
    text_w = sec.page_width - sec.left_margin - sec.right_margin

    # 1. 应无首行缩进的段落：表格单元格、题注、独占一段的图片
    cells = [p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    captions = [p for p in doc.paragraphs if p.style.name == "Caption"]
    pics = [p for p in doc.paragraphs if p._p.find(".//" + qn("w:drawing")) is not None]
    for label, paras in (("表格单元格", cells), ("题注", captions), ("图片段落", pics)):
        bad = [(p.text[:8], effective_chars(doc, p, "first")) for p in paras
               if effective_chars(doc, p, "first") != 0]
        check("%s无首行缩进（%d 段）" % (label, len(paras)), paras and not bad, bad[:3])

    # 2. 封面信息表
    cover_tbls = [e for e in cover_els if e.tag == qn("w:tbl")]
    if not cover_tbls:
        check("封面有信息表", False)
        return
    tbl = cover_tbls[0]
    look = tbl.tblPr.find(qn("w:tblLook"))
    first_row = False
    if look is not None:
        if look.get(qn("w:firstRow")) is not None:
            first_row = look.get(qn("w:firstRow")) == "1"
        else:
            first_row = bool(int(look.get(qn("w:val")) or "0", 16) & 0x20)
    check("封面信息表未启用表格样式首行格式（否则会多出横线）", not first_row)
    grid = [int(g.get(qn("w:w"))) * 635 for g in tbl.iter(qn("w:gridCol"))]
    check("封面信息表总宽不超过版心", sum(grid) <= text_w + 635,
          "%.1f / %.1f cm" % (sum(grid) / 360000, text_w / 360000))
    at_max = sum(grid) >= text_w - 635
    overflow = [text_of(tc) for tr in tbl.iter(qn("w:tr"))
                for tc, w in zip(tr.findall(qn("w:tc")), grid)
                if text_len_cm(text_of(tc), 14) + 0.5 > w / 360000 and not at_max]
    check("封面字段一行放得下（放不下时列宽已到版心上限）", not overflow, overflow)

    # 3. 正文列表：编号占首行 2 字符、每级左缩进 2 字符、编号后接空格
    body_lists = [p for p in doc.paragraphs
                  if p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None
                  and not p.style.name.startswith("Heading")]
    bad = []
    for p in body_lists:
        ilvl = p._p.pPr.find(qn("w:numPr")).find(qn("w:ilvl"))
        level = int(ilvl.get(qn("w:val"))) if ilvl is not None else 0
        first, left = effective_chars(doc, p, "first"), effective_chars(doc, p, "left")
        if first != 2 or left != level * 2:
            bad.append((p.text[:8], first, left))
    check("正文列表缩进与正文段落一致（%d 项）" % len(body_lists), body_lists and not bad, bad[:3])
    numbering = doc.part.numbering_part.element
    used = {p._p.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val")) for p in body_lists}
    abs_ids = {n.find(qn("w:abstractNumId")).get(qn("w:val")) for n in numbering.findall(qn("w:num"))
               if n.get(qn("w:numId")) in used}
    tabs = sorted({a.get(qn("w:abstractNumId")) for a in numbering.findall(qn("w:abstractNum"))
                   if a.get(qn("w:abstractNumId")) in abs_ids
                   for lvl in a.findall(qn("w:lvl"))
                   if lvl.find(qn("w:suff")) is None or lvl.find(qn("w:suff")).get(qn("w:val")) != "space"})
    check("列表编号后接空格（不用制表符）", used and not tabs, tabs)

    # 4. 图片不超版心
    shapes = list(doc.inline_shapes)
    check("图片宽度不超过版心（%d 张）" % len(shapes), shapes and all(s.width <= text_w for s in shapes))

    # 5. 图题在图下、表题在表上
    body = list(doc.element.body.iterchildren())
    misplaced = []
    for i, el in enumerate(body):
        if el.tag == qn("w:p") and el.find(".//" + qn("w:drawing")) is not None:
            nxt = body[i + 1] if i + 1 < len(body) else None
            if nxt is None or not text_of(nxt).strip().startswith("图"):
                misplaced.append("图后：" + (text_of(nxt)[:10] if nxt is not None else "无"))
        if el.tag == qn("w:tbl") and el not in cover_tbls and len(el.findall(qn("w:tr"))) > 1:
            prev = body[i - 1] if i else None
            if prev is None or not text_of(prev).strip().startswith("表"):
                misplaced.append("表前：" + (text_of(prev)[:10] if prev is not None else "无"))
    check("图题在图下、表题在表上", not misplaced, misplaced[:3])


if __name__ == "__main__":
    main()
