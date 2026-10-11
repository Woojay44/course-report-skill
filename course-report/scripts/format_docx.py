#!/usr/bin/env python3
"""中文课程报告 Word 排版脚本。

功能（按需组合，详见 SKILL.md）：
  --defaults | --black-only   二选一：默认版式 / 仅统一黑色文字
  --three-line 1,3            指定表格转三线表（序号按原稿 DOCX 中表格出现顺序，含摘要外框表格）
  --fix-abstract              摘要去灰底/外框，排成论文摘要格式
  --cover cover.json          生成封面（原稿没有封面时使用）
  --existing-cover            原稿已有封面：不生成，只把它作为封面单独分节（无页码）
  --keep-front                生成封面时不删除原稿前置区的旧封面段落
  --toc                       在封面后插入自动目录（TOC 域），并渲染回填页码
  --page-numbers [roman|continuous]
                              页脚居中页码；封面无页码。
                              roman：目录/摘要用罗马数字，正文从 1 开始（默认）
                              continuous：封面计 1 页但不显示，其后连续编号
  --no-fill-toc               不做"渲染后回填目录页码"，改为让 Word 打开时更新域
  --timezone Asia/Taipei      cover.json 里值为 "auto" 的日期按该时区取当天（默认 Asia/Taipei）
  --three-line all            所有表格转三线表
"""

import argparse
import copy
import json
import re
import shutil
import subprocess
import tempfile
import unicodedata
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

BLACK = RGBColor(0, 0, 0)
LATIN = "Times New Roman"
THEME_ATTRS = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")
# 这些字体的文字保持原样（符号、代码、公式字体）
KEEP_FONTS = {
    "symbol", "wingdings", "wingdings 2", "wingdings 3", "webdings", "consolas",
    "courier new", "menlo", "monaco", "cambria math", "lucida console",
    "sf mono", "source code pro", "dejavu sans mono",
}
MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
ABS_TITLE = re.compile(r"^(摘要|abstract)$", re.I)
KW_RE = re.compile(r"^(关键词|关键字|keywords?)[:：]", re.I)
TOC_TITLES = {"目录", "目次", "contents", "tableofcontents"}
REF_TITLES = {"参考文献", "参考资料", "references", "reference", "bibliography"}
LEADER = re.compile(r"(\.\s*){4,}|[…·]{2,}")
LOG = []


def log(msg):
    LOG.append(msg)
    print("[format_docx] " + msg)



# --------------------------------------------------------------------------
# 样式规格：字体（中文）、字号、加粗、对齐、段前/段后(pt)、首行缩进(字符)、左缩进(字符)、行距、与下段同页
# --------------------------------------------------------------------------
def sp(cjk, size, bold=False, align=A.LEFT, before=0, after=0, first=0, left=0,
       line=1.5, keep=False):
    return dict(cjk=cjk, size=size, bold=bold, align=align, before=before,
                after=after, first=first, left=left, line=line, keep=keep)


SPEC = {
    "Normal": sp("宋体", 12, align=A.JUSTIFY, before=6, first=2),
    "Cover Info": sp("宋体", 12, align=A.CENTER, after=6),
    "Title": sp("黑体", 18, True, A.CENTER),
    "Heading 1": sp("黑体", 16, True, A.CENTER, before=18, after=12, keep=True),
    "Heading 2": sp("黑体", 14, True, A.LEFT, before=12, after=6, keep=True),
    "Heading 3": sp("黑体", 12, True, A.LEFT, before=12, after=6, keep=True),
    "Caption": sp("宋体", 10.5, align=A.CENTER),
    "Table Text": sp("宋体", 10.5, line=1.25),
}
EXTRA = {
    "Cover Org": sp("黑体", 16, True, A.CENTER),
    "Cover Title": sp("黑体", 22, True, A.CENTER),
    "Cover Subtitle": sp("黑体", 16, False, A.CENTER),
    "Cover Field": sp("宋体", 14, align=A.LEFT, line=1.0),
    "TOC Title": sp("黑体", 16, True, A.CENTER, after=18),
    "toc 1": sp("宋体", 12, True, A.LEFT, line=1.75),
    "toc 2": sp("宋体", 12, False, A.LEFT, line=1.75, left=2),
    "toc 3": sp("宋体", 12, False, A.LEFT, line=1.75, left=4),
    "Abstract Title": sp("黑体", 16, True, A.CENTER, after=12),
    "Abstract Text": sp("宋体", 12, align=A.JUSTIFY, first=2),
    "Keywords": sp("宋体", 12, align=A.LEFT, first=2, before=6),
    "Footer": sp("宋体", 10.5, align=A.CENTER, line=1.0),
}
ALL_SPEC = {**SPEC, **EXTRA}


# --------------------------------------------------------------------------
# 底层 XML 工具
# --------------------------------------------------------------------------
def insert_ordered(parent, child, successors):
    """按 schema 顺序插入：放到第一个 successors 之前，没有则追加。"""
    for tag in successors:
        found = parent.find(qn(tag))
        if found is not None:
            found.addprevious(child)
            return child

    parent.append(child)
    return child


def strip_theme(rfonts):
    for attr in THEME_ATTRS:
        rfonts.attrib.pop(qn("w:" + attr), None)


def set_fonts(font, element, cjk, size, bold, style_level=False, latin=LATIN):
    font.name = latin
    font.size = Pt(size)
    font.color.rgb = BLACK
    if style_level:
        font.bold = bool(bold)
    elif bold:
        font.bold = True
    rfonts = element.get_or_add_rPr().get_or_add_rFonts()
    strip_theme(rfonts)  # 主题字体属性优先级高于 eastAsia，必须清掉
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:cs"), latin)


def set_space_before(pf, pt):
    spacing = pf._element.get_or_add_pPr().get_or_add_spacing()
    for key in ("before", "beforeLines", "beforeAutospacing"):
        spacing.attrib.pop(qn("w:" + key), None)
    spacing.set(qn("w:before"), str(int(round(pt * 20))))


def set_space_after(pf, pt):
    spacing = pf._element.get_or_add_pPr().get_or_add_spacing()

    for key in ("after", "afterLines", "afterAutospacing"):
        spacing.attrib.pop(qn("w:" + key), None)
    spacing.set(qn("w:after"), str(int(round(pt * 20))))


def set_first_line(ppr, chars):
    """首行缩进按"字符"设置（随字号变化），0 表示无缩进。"""
    ind = ppr.get_or_add_ind()
    for key in ("firstLine", "firstLineChars", "hanging", "hangingChars"):
        ind.attrib.pop(qn("w:" + key), None)
    if chars:
        ind.set(qn("w:firstLineChars"), str(int(chars * 100)))
        ind.set(qn("w:firstLine"), str(int(chars * 240)))
    else:
        # 必须同时写 firstLineChars=0：Word 中字符单位优先，只写 firstLine=0
        # 挡不住从 Normal 继承来的 firstLineChars=200
        ind.set(qn("w:firstLineChars"), "0")
        ind.set(qn("w:firstLine"), "0")


def set_left(ppr, chars):
    ind = ppr.get_or_add_ind()
    for key in ("left", "leftChars", "start", "startChars"):
        ind.attrib.pop(qn("w:" + key), None)
    if chars:
        ind.set(qn("w:leftChars"), str(int(chars * 100)))
        ind.set(qn("w:left"), str(int(chars * 240)))


def set_list_indent(ppr, chars):
    """正文列表项：左缩进 = 层级 × 2 字符，首行缩进 chars 字符（覆盖编号定义里的悬挂缩进）。返回 numId。"""
    num_pr = ppr.find(qn("w:numPr"))
    ilvl = num_pr.find(qn("w:ilvl"))
    level = int(ilvl.get(qn("w:val"))) if ilvl is not None else 0
    ind = ppr.get_or_add_ind()
    for key in list(ind.attrib):
        del ind.attrib[key]
    ind.set(qn("w:leftChars"), str(level * 200))
    ind.set(qn("w:left"), str(level * 480))
    ind.set(qn("w:firstLineChars"), str(int(chars * 100)))
    ind.set(qn("w:firstLine"), str(int(chars * 240)))
    num_id = num_pr.find(qn("w:numId"))
    return num_id.get(qn("w:val")) if num_id is not None else None


def list_suffix_space(doc, num_ids):
    """编号后用空格而不是制表符：首行缩进下制表符会跳到下一个制表位，编号变长时空出一大段。"""
    if not num_ids:
        return
    numbering = doc.part.numbering_part.element
    abstract_ids = set()
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) in num_ids:
            ref = num.find(qn("w:abstractNumId"))
            if ref is not None:
                abstract_ids.add(ref.get(qn("w:val")))
    for absnum in numbering.findall(qn("w:abstractNum")):
        if absnum.get(qn("w:abstractNumId")) not in abstract_ids:
            continue
        for lvl in absnum.findall(qn("w:lvl")):
            suff = lvl.find(qn("w:suff"))
            if suff is None:
                suff = OxmlElement("w:suff")
                insert_ordered(lvl, suff, ["w:lvlText", "w:lvlPicBulletId", "w:legacy",
                                           "w:lvlJc", "w:pPr", "w:rPr"])
            suff.set(qn("w:val"), "space")


def set_hanging(ppr, chars):
    ind = ppr.get_or_add_ind()
    for key in ("firstLine", "firstLineChars", "hanging", "hangingChars",
                "left", "leftChars", "start", "startChars"):
        ind.attrib.pop(qn("w:" + key), None)
    ind.set(qn("w:leftChars"), str(int(chars * 100)))
    ind.set(qn("w:left"), str(int(chars * 240)))
    ind.set(qn("w:hangingChars"), str(int(chars * 100)))
    ind.set(qn("w:hanging"), str(int(chars * 240)))


RPR_AFTER_SPACING = ["w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u",
                     "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl",
                     "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath"]


def set_char_spacing(run_el, twips):
    rpr = run_el.get_or_add_rPr()
    old = rpr.find(qn("w:spacing"))
    if old is not None:
        rpr.remove(old)
    el = OxmlElement("w:spacing")
    el.set(qn("w:val"), str(twips))
    insert_ordered(rpr, el, RPR_AFTER_SPACING)



def text_of(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def norm(s):
    return re.sub(r"\s+", "", s or "")


def has_drawing(el):
    return any(True for _ in el.iter(qn("w:drawing"), qn("w:pict"), qn("w:object")))


def has_math(el):
    return any(True for _ in el.iter(MATH + "oMath", MATH + "oMathPara"))


def is_blank_like(el):
    """没有文字、图片、公式的段落。"""
    return (el.tag == qn("w:p") and not text_of(el).strip()
            and not has_drawing(el) and not has_math(el))


def has_struct_break(el):
    ppr = el.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
        return True
    return any(br.get(qn("w:type")) == "page" for br in el.iter(qn("w:br")))


# 书签等"透明"元素：pandoc 会在标题前后放 bookmarkStart/End，它们不是内容
TRANSPARENT = {qn("w:bookmarkStart"), qn("w:bookmarkEnd"), qn("w:proofErr"),
               qn("w:permStart"), qn("w:permEnd")}


def prev_real(el):
    prev = el.getprevious()
    while prev is not None and prev.tag in TRANSPARENT:
        prev = prev.getprevious()
    return prev


def pstyle(el, doc):
    try:
        return Paragraph(el, doc._body).style.name
    except Exception:
        return ""


def heading_level(el, doc):
    if el.tag != qn("w:p"):
        return 0
    m = re.match(r"^Heading (\d)$", pstyle(el, doc))
    return int(m.group(1)) if m else 0


def shrink_paragraph(p_el):
    """把承载分节符/域结束符的段落压成 1pt 高，避免多出空白页。"""
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p_el.insert(0, ppr)
    for old in ppr.findall(qn("w:spacing")):
        ppr.remove(old)

    spacing = OxmlElement("w:spacing")
    for key, val in (("before", "0"), ("after", "0"), ("line", "20"), ("lineRule", "exact")):
        spacing.set(qn("w:" + key), val)
    insert_ordered(ppr, spacing, ["w:ind", "w:contextualSpacing", "w:mirrorIndents",
                                  "w:suppressOverlap", "w:jc", "w:textDirection",
                                  "w:textAlignment", "w:textboxTightWrap",
                                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
                                  "w:sectPr", "w:pPrChange"])


# --------------------------------------------------------------------------
# 样式定义与段落格式
# --------------------------------------------------------------------------
def text_width_cm(doc):
    sec = doc.sections[-1]
    if None in (sec.page_width, sec.left_margin, sec.right_margin):
        return 16.0
    return (sec.page_width - sec.left_margin - sec.right_margin) / 360000


def text_len_cm(text, size_pt):
    """估算单行文字宽度：全角字符按 1 个字号宽，半角按 0.55。"""
    ems = sum(1.0 if unicodedata.east_asian_width(ch) in "WF" else 0.55 for ch in text)
    return ems * size_pt / 72 * 2.54


def cover_col_widths(doc, labels, values, size_pt):
    """封面信息表列宽：默认 3.8 + 7.2 cm，内容更长时加宽，总宽不超过版心。"""
    pad = 0.5  # 单元格左右边距
    label_w = max(3.8, max(text_len_cm(t, size_pt) for t in labels) + pad)
    value_w = max(7.2, max((text_len_cm(t, size_pt) for t in values), default=0) + pad + 0.6)
    value_w = min(value_w, text_width_cm(doc) - label_w)
    return Cm(label_w), Cm(value_w)


def set_table_width(table, width):
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        insert_ordered(tblpr, tblw, ["w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders",
                                     "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"])
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(int(width / 635)))  # EMU → twips


def plain_table_look(table):
    """关闭表格样式的首行/首列等条件格式（pandoc 默认表格样式给首行加了下边框）。"""
    tblpr = table._tbl.tblPr
    look = tblpr.find(qn("w:tblLook"))
    if look is None:
        look = OxmlElement("w:tblLook")
        tblpr.append(look)
    look.set(qn("w:val"), "0600")
    for key, val in (("firstRow", "0"), ("lastRow", "0"), ("firstColumn", "0"),
                     ("lastColumn", "0"), ("noHBand", "1"), ("noVBand", "1")):
        look.set(qn("w:" + key), val)


def style_format(doc, style, name):
    s = ALL_SPEC[name]
    set_fonts(style.font, style._element, s["cjk"], s["size"], s["bold"], style_level=True)
    ppr = style._element.get_or_add_pPr()
    if name == "Title":
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)
    pf = style.paragraph_format
    pf.alignment = s["align"]
    pf.line_spacing = s["line"]
    set_space_before(pf, s["before"])
    set_space_after(pf, s["after"])
    set_first_line(ppr, s["first"])
    set_left(ppr, s["left"])

    pf.keep_with_next = s["keep"]
    pf.widow_control = True
    if name.startswith("toc "):
        old_tabs = ppr.find(qn("w:tabs"))
        if old_tabs is not None:
            ppr.remove(old_tabs)  # 重复调用时不叠加制表位
        pf.keep_together = False
        pf.page_break_before = False  # 目录条目不得带分页属性
        pf.tab_stops.add_tab_stop(Cm(text_width_cm(doc)), WD_TAB_ALIGNMENT.RIGHT,
                                  WD_TAB_LEADER.DOTS)


def define_styles(doc, names):
    for name in names:
        if name not in doc.styles:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = doc.styles["Normal"]
            if name.startswith("toc "):
                style.element.set(qn("w:styleId"), "TOC" + name[-1])
        style_format(doc, doc.styles[name], name)


def fix_doc_defaults(doc):
    dd = doc.styles.element.find(qn("w:docDefaults"))
    if dd is None:
        return
    for rf in dd.iter(qn("w:rFonts")):
        strip_theme(rf)
        rf.set(qn("w:ascii"), LATIN)
        rf.set(qn("w:hAnsi"), LATIN)
        rf.set(qn("w:cs"), LATIN)
        rf.set(qn("w:eastAsia"), "宋体")



STYLE_MAP = {
    "Body Text": "Normal", "First Paragraph": "Normal", "Compact": "Normal",
    "Captioned Figure": "Normal", "Figure": "Normal", "Normal (Web)": "Normal",
    "Table Caption": "Caption", "Image Caption": "Caption",
}


def normalize_styles(doc):
    """pandoc 等工具生成的 docx 使用 Body Text/Compact 等样式，统一并入 Normal/Caption。"""
    count = 0
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        paragraphs.extend(iter_table_paragraphs(table))
    for p in paragraphs:
        target = STYLE_MAP.get(p.style.name)
        if target and target in doc.styles:
            p.style = doc.styles[target]
            count += 1
    if count:
        log("已把 %d 个 pandoc 样式段落并入 Normal/Caption" % count)


def tidy_code_style(doc):
    """pandoc 的 Source Code 样式基于 Normal，会继承两端对齐和首行缩进，这里改回左对齐、无缩进。"""
    if "Source Code" in doc.styles:
        style = doc.styles["Source Code"]
        ppr = style._element.get_or_add_pPr()
        style.paragraph_format.alignment = A.LEFT
        style.paragraph_format.line_spacing = 1.15
        set_first_line(ppr, 0)
        set_left(ppr, 0)
        set_space_before(style.paragraph_format, 3)
        set_space_after(style.paragraph_format, 3)


def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for inner in cell.tables:
                yield from iter_table_paragraphs(inner)


def run_keeps_font(run):
    rpr = run._element.find(qn("w:rPr"))
    if run._element.find(qn("w:sym")) is not None:
        return True

    if rpr is not None:
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            for key in ("ascii", "hAnsi", "eastAsia"):
                if (rf.get(qn("w:" + key)) or "").lower() in KEEP_FONTS:
                    return True
    return False


def apply_paragraph(paragraph, names, gen):
    name = paragraph.style.name
    p = paragraph._p
    if name not in names or id(p) in gen:
        return
    ppr = p.find(qn("w:pPr"))
    if ppr is not None:
        spacing = ppr.find(qn("w:spacing"))
        if spacing is not None and spacing.get(qn("w:lineRule")) == "exact":
            return  # 承载分节符的 1pt 段落
    if has_math(p) and not text_of(p).strip():
        return  # 独占一段的公式保持 Word 原生对齐；行内公式的段落照常排版
    s = ALL_SPEC[name]
    pf = paragraph.paragraph_format
    has_pic = has_drawing(p)
    has_num = ppr is not None and ppr.find(qn("w:numPr")) is not None
    explicit_jc = ppr is not None and ppr.find(qn("w:jc")) is not None
    if not (name == "Table Text" and explicit_jc):
        pf.alignment = s["align"]
    pf.line_spacing = s["line"]
    set_space_before(pf, s["before"])
    set_space_after(pf, s["after"])
    num_id = None
    if has_num:
        # 带自动编号的标题保留编号自带的缩进；正文列表按普通段落排：编号占首行缩进位置，折行顶格
        if not name.startswith("Heading"):
            num_id = set_list_indent(ppr, s["first"])
    else:
        set_first_line(p.get_or_add_pPr(), 0 if has_pic else s["first"])
    if s["keep"]:
        pf.keep_with_next = True
    if name == "Caption" and re.match(r"^表\s*\d", text_of(p).strip()):
        pf.keep_with_next = True  # 表题在表上，必须和表同页
    for run in paragraph.runs:
        if run_keeps_font(run):
            run.font.color.rgb = BLACK
            continue
        cjk = s["cjk"]
        if name == "Keywords" and run.bold:
            cjk = "黑体"
        set_fonts(run.font, run._element, cjk, s["size"], s["bold"])
    return num_id


def apply_all(doc, names, gen, defaults):
    list_ids = set()
    for paragraph in doc.paragraphs:
        list_ids.add(apply_paragraph(paragraph, names, gen))
    for table in doc.tables:
        for paragraph in iter_table_paragraphs(table):
            if defaults and id(paragraph._p) not in gen and paragraph.style.name == "Normal":
                paragraph.style = doc.styles["Table Text"]
            list_ids.add(apply_paragraph(paragraph, names, gen))
    list_ids.discard(None)
    list_suffix_space(doc, list_ids)


# --------------------------------------------------------------------------
# 颜色、表格、图片、题注
# --------------------------------------------------------------------------
def force_black(doc):
    for style in doc.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            style.font.color.rgb = BLACK
    for part in doc.part.package.parts:
        element = getattr(part, "element", None)
        if element is None:
            continue
        for run in element.iter(qn("w:r")):
            rpr = run.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                run.insert(0, rpr)
            color = rpr.find(qn("w:color"))
            if color is None:
                color = OxmlElement("w:color")
                insert_ordered(rpr, color, ["w:spacing"] + RPR_AFTER_SPACING)
            color.set(qn("w:val"), "000000")
            for key in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn("w:" + key), None)


BORDER_ORDER = ["top", "left", "start", "bottom", "right", "end", "insideH", "insideV",
                "tl2br", "tr2bl"]



def set_border(parent, edge, size=None):
    """size 为 1/8 pt；None 表示无边框。按 schema 顺序插入。"""
    is_cell = parent.tag == qn("w:tcPr")
    borders_name = "w:tcBorders" if is_cell else "w:tblBorders"
    borders = parent.find(qn(borders_name))
    if borders is None:
        borders = OxmlElement(borders_name)
        if is_cell:
            insert_ordered(parent, borders, ["w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
                                             "w:tcFitText", "w:vAlign", "w:hideMark"])
        else:
            insert_ordered(parent, borders, ["w:shd", "w:tblLayout", "w:tblCellMar",
                                             "w:tblLook", "w:tblCaption", "w:tblDescription"])
    border = borders.find(qn("w:" + edge))
    if border is None:
        border = OxmlElement("w:" + edge)
        rank = BORDER_ORDER.index(edge)
        placed = False
        for other in borders:
            name = other.tag.split("}")[1]
            if name in BORDER_ORDER and BORDER_ORDER.index(name) > rank:
                other.addprevious(border)
                placed = True
                break
        if not placed:
            borders.append(border)
    border.set(qn("w:val"), "single" if size else "nil")
    if size:
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:space"), "0")

        border.set(qn("w:color"), "000000")
    else:
        for key in ("sz", "space", "color"):
            border.attrib.pop(qn("w:" + key), None)


def three_line_table(table):
    for edge in ("left", "right", "insideH", "insideV"):
        set_border(table._tbl.tblPr, edge)
    for edge in ("top", "bottom"):
        set_border(table._tbl.tblPr, edge, 12)
    for row in table.rows:
        for cell in row.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            old = tcpr.find(qn("w:tcBorders"))
            if old is not None:
                tcpr.remove(old)
    if table.rows:
        for cell in table.rows[0].cells:
            set_border(cell._tc.get_or_add_tcPr(), "bottom", 6)


def tidy_tables(doc):
    """表格居中；表头行跨页重复；行不跨页断开。"""
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            if trpr.find(qn("w:cantSplit")) is None:
                insert_ordered(trpr, OxmlElement("w:cantSplit"),
                               ["w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden"])
            if i == 0 and trpr.find(qn("w:tblHeader")) is None:
                insert_ordered(trpr, OxmlElement("w:tblHeader"),
                               ["w:tblCellSpacing", "w:jc", "w:hidden"])


def tidy_figures(doc, gen):
    """图片缩放到版心内、单独成段的图片居中、图与图题同页。"""
    max_w = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
    for shape in doc.inline_shapes:
        if shape.width > max_w:
            ratio = max_w / shape.width
            shape.height = int(shape.height * ratio)

            shape.width = int(max_w)
    for paragraph in doc.paragraphs:
        p = paragraph._p
        if id(p) in gen or not has_drawing(p) or text_of(p).strip():
            continue
        paragraph.alignment = A.CENTER
        set_first_line(p.get_or_add_pPr(), 0)
        paragraph.paragraph_format.keep_with_next = True


def fix_caption_positions(doc):
    """图题放图下、表题放表上（只移动样式为 Caption 的题注，条件严格）。"""
    body = doc.element.body

    def is_caption(el):
        return el is not None and el.tag == qn("w:p") and pstyle(el, doc) == "Caption"

    moved = 0
    for el in list(body.iterchildren()):
        if not is_caption(el):
            continue
        t = text_of(el).strip()
        prev, nxt = el.getprevious(), el.getnext()
        if re.match(r"^表\s*\d", t) and prev is not None and prev.tag == qn("w:tbl"):
            before = prev.getprevious()
            if not (is_caption(before) and re.match(r"^表\s*\d", text_of(before).strip())):
                prev.addprevious(el)
                moved += 1
        elif (re.match(r"^图\s*\d", t) and nxt is not None and nxt.tag == qn("w:p")
              and has_drawing(nxt) and not text_of(nxt).strip()):
            before = el.getprevious()
            if not (before is not None and before.tag == qn("w:p") and has_drawing(before)):

                nxt.addnext(el)
                moved += 1
    if moved:
        log("已调整 %d 个题注位置（图题在图下、表题在表上）" % moved)


def format_references(doc):
    """参考文献条目：悬挂缩进 2 字符（跳过自动编号列表）。"""
    body = doc.element.body
    active = False
    level = 0
    count = 0
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        lvl = heading_level(el, doc)
        if lvl:
            if norm(text_of(el)).lower() in REF_TITLES:
                active, level = True, lvl
                continue
            if active and lvl <= level:
                active = False
        if active and pstyle(el, doc) == "Normal" and text_of(el).strip():
            ppr = el.get_or_add_pPr()
            if ppr.find(qn("w:numPr")) is None:
                set_hanging(ppr, 2)
                Paragraph(el, doc._body).alignment = A.LEFT
                count += 1
    if count:
        log("参考文献 %d 条已设为悬挂缩进 2 字符" % count)


# --------------------------------------------------------------------------
# 摘要：拆外框、去底纹、排成论文摘要
# --------------------------------------------------------------------------
def is_abs_marker(el):
    if el.tag != qn("w:p"):
        return False
    t = norm(text_of(el))
    return bool((ABS_TITLE.match(t) and len(t) <= 9) or KW_RE.match(t))


def unwrap_containers(doc):
    """摘要被放在单列表格或文本框里（常带灰底）时，把内容还原成普通段落。"""
    body = doc.element.body
    for el in list(body.iterchildren()):
        if el.tag == qn("w:tbl"):
            rows = el.findall(qn("w:tr"))

            cells = [tc for tr in rows for tc in tr.findall(qn("w:tc"))]
            if cells and len(cells) == len(rows):
                inner = [c for tc in cells for c in tc if c.tag in (qn("w:p"), qn("w:tbl"))]
                if any(is_abs_marker(c) for c in inner):
                    for c in inner:
                        el.addprevious(c)
                    body.remove(el)
                    log("已拆除摘要外框表格")
        elif el.tag == qn("w:p"):
            boxes = list(el.iter(qn("w:txbxContent")))
            if boxes:
                inner = [c for c in boxes[0] if c.tag == qn("w:p")]
                if any(is_abs_marker(c) for c in inner):
                    for c in inner:
                        el.addprevious(c)
                    for run in el.findall(qn("w:r")):
                        if run.find(".//" + qn("w:txbxContent")) is not None:
                            el.remove(run)
                    if is_blank_like(el):
                        body.remove(el)
                    log("已拆除摘要文本框")


def find_abstract_blocks(doc, cap=40):
    body = doc.element.body
    children = [e for e in body.iterchildren() if e.tag != qn("w:sectPr")]
    blocks, i = [], 0
    while i < len(children):
        el = children[i]
        t = norm(text_of(el)) if el.tag == qn("w:p") else ""
        if el.tag == qn("w:p") and ABS_TITLE.match(t) and len(t) <= 9:
            els, j = [el], i + 1

            while j < len(children) and len(els) < cap:
                e = children[j]
                if e.tag in TRANSPARENT:
                    j += 1
                    continue
                if e.tag != qn("w:p"):
                    break
                tt = norm(text_of(e))
                if (ABS_TITLE.match(tt) and len(tt) <= 9) or heading_level(e, doc):
                    break
                els.append(e)
                j += 1
                if KW_RE.match(tt):
                    break
            blocks.append({"title": el, "els": els})
            i = j
        else:
            i += 1
    return blocks


PPR_KEEP = {qn("w:pStyle"), qn("w:sectPr"), qn("w:rPr")}


def clean_paragraph(el):
    """清除底纹、边框、框架、缩进等直接格式，交给样式统一控制。"""
    ppr = el.find(qn("w:pPr"))
    if ppr is not None:
        for child in list(ppr):
            if child.tag not in PPR_KEEP:
                ppr.remove(child)
        mark = ppr.find(qn("w:rPr"))
        if mark is not None:
            for tag in ("w:shd", "w:highlight", "w:bdr"):
                for x in mark.findall(qn(tag)):
                    mark.remove(x)
    for run in el.iter(qn("w:r")):
        rpr = run.find(qn("w:rPr"))
        if rpr is not None:
            for tag in ("w:shd", "w:highlight", "w:bdr"):
                for x in rpr.findall(qn(tag)):
                    mark.remove(x)



def rebuild_keywords(doc, el):
    full = text_of(el)
    m = re.match(r"^\s*(关键词|关键字|key\s*words?)\s*([:：])(\s*)(.*)$", full, re.I | re.S)
    if not m:
        return
    for child in list(el):
        if child.tag != qn("w:pPr"):
            el.remove(child)
    p = Paragraph(el, doc._body)
    p.add_run(m.group(1) + m.group(2) + m.group(3)).bold = True  # 保留冒号后原有空格
    p.add_run(m.group(4))


def fix_abstract(doc):
    unwrap_containers(doc)
    define_styles(doc, ["Abstract Title", "Abstract Text", "Keywords"])
    blocks = find_abstract_blocks(doc)
    if not blocks:
        log("未找到\"摘要/Abstract\"标题，跳过摘要处理")
        return
    body = doc.element.body
    for blk in blocks:
        alive = []
        for el in blk["els"]:
            if el is not blk["title"] and is_blank_like(el):
                body.remove(el)
                continue
            clean_paragraph(el)
            alive.append(el)
        for el in alive:
            p = Paragraph(el, doc._body)
            t = norm(text_of(el))

            if el is blk["title"]:
                p.style = doc.styles["Abstract Title"]
                p.paragraph_format.page_break_before = True
                if t == "摘要":
                    for run in p.runs:
                        set_char_spacing(run._element, 320)  # "摘  要"，不改动文字本身
            elif KW_RE.match(t):
                p.style = doc.styles["Keywords"]
                rebuild_keywords(doc, el)
            else:
                p.style = doc.styles["Abstract Text"]
        blk["els"] = alive
    log("摘要已处理 %d 处（去底纹/外框，套摘要样式，每处另起一页）" % len(blocks))


# --------------------------------------------------------------------------
# 旧目录清理、前置区清理
# --------------------------------------------------------------------------
def remove_old_toc(doc):
    body = doc.element.body
    removed = 0
    for el in list(body.iterchildren()):
        if el.tag == qn("w:sdt"):
            gallery = [g.get(qn("w:val")) or "" for g in el.iter(qn("w:docPartGallery"))]
            toc_field = any(re.match(r"\s*TOC\b", it.text or "") for it in el.iter(qn("w:instrText")))
            if any("Table of Contents" in g for g in gallery) or toc_field:
                body.remove(el)
                removed += 1
    in_toc, depth = False, 0
    for el in list(body.iterchildren()):
        if el.tag != qn("w:p"):
            continue
        style = pstyle(el, doc).lower()
        instr = [it.text or "" for it in el.iter(qn("w:instrText"))]
        simple = [f.get(qn("w:instr")) or "" for f in el.iter(qn("w:fldSimple"))]
        starts = any(re.match(r"\s*TOC\b", x) for x in instr + simple)
        if starts and not in_toc:
            if any(re.match(r"\s*TOC\b", x) for x in instr):
                in_toc, depth = True, 0
        if in_toc or starts or style.startswith("toc ") or style == "toc heading":
            if in_toc:
                for fc in el.iter(qn("w:fldChar")):
                    kind = fc.get(qn("w:fldCharType"))
                    depth += 1 if kind == "begin" else -1 if kind == "end" else 0
                if depth <= 0:
                    in_toc = False
            body.remove(el)
            removed += 1
    titles = [e for e in list(body.iterchildren())[:80]
              if e.tag == qn("w:p") and norm(text_of(e)).lower() in TOC_TITLES]
    if removed:
        for e in titles:
            body.remove(e)
        log("已删除原稿旧目录（%d 个元素）" % removed)
    elif titles:
        log("警告：发现\"目录\"标题但没有目录域，疑似手工目录，请人工确认并删除后再运行")



def strip_blank_before(el):
    prev = el.getprevious()
    while prev is not None:
        if prev.tag in TRANSPARENT:
            prev = prev.getprevious()
        elif is_blank_like(prev) and not has_struct_break(prev):
            nxt = prev.getprevious()
            prev.getparent().remove(prev)
            prev = nxt
        else:
            break


def clean_front_breaks(doc, body_start):
    body = doc.element.body
    for el in list(body.iterchildren()):
        if el is body_start:
            break
        if el.tag != qn("w:p"):
            continue
        if is_blank_like(el) and has_struct_break(el):
            body.remove(el)
            continue
        ppr = el.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            ppr.remove(ppr.find(qn("w:sectPr")))
        for br in list(el.iter(qn("w:br"))):
            if br.get(qn("w:type")) == "page":
                br.getparent().remove(br)



def find_body_start(doc, blocks):
    body = doc.element.body
    if blocks:
        e = blocks[-1]["els"][-1].getnext()
        while e is not None:
            nxt = e.getnext()
            if e.tag in TRANSPARENT:
                pass
            elif e.tag == qn("w:p") and is_blank_like(e):
                body.remove(e)
            else:
                break
            e = nxt
        return e if (e is not None and e.tag != qn("w:sectPr")) else None
    for e in body.iterchildren():
        if e.tag == qn("w:p") and heading_level(e, doc):
            return e
    return None


def drop_front(doc, anchor, gen):
    """生成新封面前，删除锚点之前的旧封面短段落/小表格，
并打印被删内容供核对。"""
    body = doc.element.body
    dropped = []
    for el in list(body.iterchildren()):
        if el is anchor or el.tag == qn("w:sectPr"):
            break
        if el.tag == qn("w:p"):
            t = text_of(el).strip()
            if has_drawing(el):
                log("警告：前置区含图片，已保留，请确认是否属于旧封面")
            elif len(t) <= 80:
                if t:
                    dropped.append(t)
                body.remove(el)
            else:
                log("警告：前置区有较长段落（%d 字），已保留：%s…" % (len(t), t[:20]))
        elif el.tag == qn("w:tbl"):
            t = " | ".join(text_of(tc).strip() for tc in el.iter(qn("w:tc")))
            if len(t) <= 200:
                dropped.append("[表格] " + t)
                body.remove(el)
    if dropped:
        log("已删除旧封面内容（请核对是否已写入 cover.json）：" + " / ".join(dropped))


# --------------------------------------------------------------------------
# 封面、目录、分节、页码
# --------------------------------------------------------------------------
def build_cover(doc, info, gen, gen_refs):
    els = []

    def track(el):

        gen.add(id(el))
        gen_refs.append(el)

    def add(text, style, before=0, after=0):
        p = doc.add_paragraph(text, style=style)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        track(p._p)
        els.append(p._p)
        return p

    if info.get("org"):
        add(info["org"], "Cover Org", after=12)
    if info.get("course"):
        add(info["course"], "Cover Subtitle", after=0)
    add(info["title"], "Cover Title", before=90, after=10 if info.get("subtitle") else 96)
    if info.get("subtitle"):
        add(info["subtitle"], "Cover Subtitle", after=96)

    rows = info.get("fields") or []
    if rows:
        labels = []
        for label, _ in rows:
            label = str(label).strip()
            if len(label) == 2:
                label = label[0] + "　　" + label[1]
            if not label.endswith(("：", ":")):
                label += "："
            labels.append(label)
        widths = cover_col_widths(doc, labels, [str(v or "") for _, v in rows],
                                  ALL_SPEC["Cover Field"]["size"])
        table = doc.add_table(rows=len(rows), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for col, w in zip(table.columns, widths):
            col.width = w
        set_table_width(table, sum(widths))
        plain_table_look(table)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            set_border(table._tbl.tblPr, edge)
        for row, label, (_, value) in zip(table.rows, labels, rows):
            row.height = Cm(1.2)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            c0, c1 = row.cells
            for cell, w in zip((c0, c1), widths):
                cell.width = w
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p0, p1 = c0.paragraphs[0], c1.paragraphs[0]
            p0.style = p1.style = doc.styles["Cover Field"]
            p0.alignment, p1.alignment = A.RIGHT, A.CENTER
            p0.add_run(label)
            p1.add_run(str(value or ""))
            set_border(c1._tc.get_or_add_tcPr(), "bottom", 6)
            track(p0._p)
            track(p1._p)
        track(table._tbl)
        els.append(table._tbl)
    if info.get("date"):
        add(info["date"], "Cover Info", before=48)
    body = doc.element.body
    for k, el in enumerate(els):
        body.insert(k, el)
    return els


def today_cn(tz_name):
    try:
        from zoneinfo import ZoneInfo
        now = datetime.now(ZoneInfo(tz_name))
    except Exception:
        now = datetime.now(timezone(timedelta(hours=8)))
        log("警告：时区 %s 不可用，已按 UTC+8 计算日期" % tz_name)
    return "%d年%d月%d日" % (now.year, now.month, now.day)


def resolve_auto(info, tz_name):
    """cover.json 中值为 auto 的日期，按时区取当天（不补零）。"""
    resolved = None
    if str(info.get("date", "")).strip().lower() == "auto":
        resolved = info["date"] = today_cn(tz_name)
    for row in info.get("fields") or []:
        if len(row) > 1 and str(row[1]).strip().lower() == "auto":
            resolved = row[1] = today_cn(tz_name)
    if resolved:
        log("封面日期按 %s 取为 %s" % (tz_name, resolved))



def fld_run(p_el, kind):
    run = OxmlElement("w:r")
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), kind)
    run.append(fc)
    p_el.append(run)


def instr_run(p_el, text):
    run = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = text
    run.append(it)
    p_el.append(run)


def collect_headings(doc, levels):
    out = []
    for el in doc.element.body.iterchildren():
        lvl = heading_level(el, doc)
        if lvl and lvl <= levels:
            text = " ".join(text_of(el).split())
            if text:
                out.append((lvl, text))
    return out


def build_toc(doc, levels, gen, gen_refs):
    headings = collect_headings(doc, levels)

    if not headings:
        raise SystemExit("未找到 Heading 1–%d 样式的标题，无法生成目录；请先给标题套用标题样式" % levels)
    define_styles(doc, ["TOC Title"] + ["toc %d" % i for i in (1, 2, 3)])
    els, entries = [], []
    title = doc.add_paragraph(style="TOC Title")
    set_char_spacing(title.add_run("目录")._element, 320)
    els.append(title._p)
    for i, (lvl, text) in enumerate(headings):
        p = doc.add_paragraph(style="toc %d" % min(lvl, 3))
        if i == 0:
            fld_run(p._p, "begin")
            instr_run(p._p, ' TOC \\o "1-%d" \\h \\z \\u ' % levels)
            fld_run(p._p, "separate")
        p.add_run(text)
        p.add_run().add_tab()
        page_run = p.add_run("1")
        entries.append({"level": lvl, "text": text, "t": page_run._element.find(qn("w:t"))})
        els.append(p._p)
    end = doc.add_paragraph()
    fld_run(end._p, "end")
    shrink_paragraph(end._p)
    els.append(end._p)
    for el in els:
        gen.add(id(el))
        gen_refs.append(el)
    return els, entries


def force_single_column(sect):
    cols = sect.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        insert_ordered(sect, cols, ["w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
                                    "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
                                    "w:printerSettings"])
    for child in list(cols):
        cols.remove(child)
    for key in list(cols.attrib):
        del cols.attrib[key]
    cols.set(qn("w:num"), "1")
    cols.set(qn("w:space"), "720")


def make_sectpr(doc):
    sect = copy.deepcopy(doc.element.body.sectPr)
    for tag in ("w:headerReference", "w:footerReference", "w:titlePg", "w:pgNumType", "w:type"):
        for x in sect.findall(qn(tag)):
            sect.remove(x)
    force_single_column(sect)  # 封面/目录/摘要分节一律单栏，不继承原稿分栏
    return sect


def set_pgnum(sect, fmt, start):
    for x in sect.findall(qn("w:pgNumType")):
        sect.remove(x)
    if fmt is None:
        return
    el = OxmlElement("w:pgNumType")
    el.set(qn("w:fmt"), fmt)
    if start:
        el.set(qn("w:start"), str(start))
    insert_ordered(sect, el, ["w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
                              "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
                              "w:printerSettings"])



def end_section_after(el, sect):
    if el.tag != qn("w:p"):
        p = OxmlElement("w:p")
        el.addnext(p)
        shrink_paragraph(p)
        el = p
    ppr = el.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        el.insert(0, ppr)
    for old in ppr.findall(qn("w:sectPr")):
        ppr.remove(old)
    ppr.append(sect)
    return el


def set_footer(doc, section, show_number):
    section.different_first_page_header_footer = False
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraphs = footer.paragraphs
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    p = paragraphs[0]
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)
    p.style = doc.styles["Footer"]
    p.alignment = A.CENTER
    if show_number:
        fld_run(p._p, "begin")
        instr_run(p._p, " PAGE ")
        fld_run(p._p, "separate")
        p.add_run("1")
        fld_run(p._p, "end")


def set_update_fields(doc):
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        insert_ordered(settings, el, ["w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr",
                                      "w:compat", "w:docVars", "w:rsids", "w:attachedSchema",
                                      "w:themeFontLang", "w:clrSchemeMapping"])


def build_structure(doc, opts, gen, gen_refs):
    """封面 → 目录 → 摘要 → 正文 的分节、页码与目录。返回 (toc_entries, body_start_text)。"""
    body = doc.element.body
    if opts.toc:
        remove_old_toc(doc)
    if opts.fix_abstract:
        fix_abstract(doc)
    blocks = find_abstract_blocks(doc)
    body_start = find_body_start(doc, blocks)
    if body_start is None:
        raise SystemExit("找不到摘要或标题，无法确定正文起点；请先给标题套用 Heading 样式")
    pre_anchor = blocks[0]["title"] if blocks else body_start
    clean_front_breaks(doc, body_start)
    ppr = body_start.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:pageBreakBefore")) is not None:
        ppr.remove(ppr.find(qn("w:pageBreakBefore")))  # 正文靠分节符另起一页
    body_start_text = text_of(body_start).strip()

    if opts.cover:
        info = json.loads(Path(opts.cover).read_text(encoding="utf-8"))
        if not info.get("title"):
            raise SystemExit("cover.json 必须包含 title")
        resolve_auto(info, opts.timezone)
        define_styles(doc, ["Cover Org", "Cover Title", "Cover Subtitle", "Cover Field", "Cover Info"])
        if not opts.keep_front:
            drop_front(doc, pre_anchor, gen)
        build_cover(doc, info, gen, gen_refs)

    entries = []
    toc_first = None
    if opts.toc:
        toc_els, entries = build_toc(doc, opts.toc_levels, gen, gen_refs)
        for el in toc_els:
            pre_anchor.addprevious(el)
        toc_first = toc_els[0]
    first_after_cover = toc_first if toc_first is not None else pre_anchor
    strip_blank_before(first_after_cover)
    has_cover = prev_real(first_after_cover) is not None
    has_front = bool(opts.toc or blocks)

    if has_cover:
        end_section_after(prev_real(first_after_cover), make_sectpr(doc))

    if has_front:
        strip_blank_before(body_start)
        end_section_after(prev_real(body_start), make_sectpr(doc))
    for s in body.iter(qn("w:sectPr")):
        for t in s.findall(qn("w:type")):
            s.remove(t)

    if opts.page_numbers:
        sections = doc.sections
        scheme = opts.page_numbers
        i = 0
        if has_cover:
            set_footer(doc, sections[i], False)
            i += 1
        if has_front:
            set_footer(doc, sections[i], True)
            if scheme == "roman":
                set_pgnum(sections[i]._sectPr, "lowerRoman", 1)
            i += 1
        set_footer(doc, sections[i], True)
        if scheme == "roman" and (has_cover or has_front):
            set_pgnum(sections[i]._sectPr, "decimal", 1)
    log("分节：封面=%s，目录/摘要=%s，页码方案=%s" % (
        "有" if has_cover else "无", "有" if has_front else "无", opts.page_numbers or "未设置"))
    return entries, body_start_text


# --------------------------------------------------------------------------
# 目录页码回填：先渲染成 PDF，再按标题所在页写回
# --------------------------------------------------------------------------
def render_pdf(docx_path, outdir):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    try:
        subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(outdir),
                        str(docx_path)], capture_output=True, timeout=240, check=True)
    except Exception:
        return None
    pdf = Path(outdir) / (Path(docx_path).stem + ".pdf")
    return pdf if pdf.exists() else None


def pdf_pages(pdf):
    if shutil.which("pdftotext"):
        try:
            out = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True,
                                 text=True, timeout=120).stdout
            pages = out.split("\f")
            if pages and not pages[-1].strip():
                pages = pages[:-1]
            return pages
        except Exception:
            pass
    try:
        from pypdf import PdfReader
        return [pg.extract_text() or "" for pg in PdfReader(str(pdf)).pages]
    except Exception:
        return None


def line_matches(page, target):
    tn = norm(target)
    for line in page.splitlines():
        if LEADER.search(line):
            continue  # 目录行（带点线）不算
        ln = norm(line)
        if ln == tn or (len(ln) >= 6 and tn.startswith(ln)):
            return True
    return False


def find_page(pages, target, start):
    for idx in range(start, len(pages)):
        if line_matches(pages[idx], target):
            return idx
    return None



def body_columns(doc):
    """返回正文节的栏数（最后一个 sectPr），默认 1。"""
    cols = doc.element.body.sectPr.find(qn("w:cols"))
    if cols is None:
        return 1
    return int(cols.get(qn("w:num")) or 1)


def body_is_multicol(doc):
    """文档最后一节（正文节）是否为多栏排版。"""
    cols = doc.element.body.sectPr.find(qn("w:cols"))
    if cols is None:
        return False
    try:
        return int(cols.get(qn("w:num"), "1")) > 1
    except ValueError:
        return False


def check_toc_layout(pdf, pages, first):
    """目录跨页时，非末页的页尾不应大片留白（提前断页的征兆）。"""
    toc_idx = [i for i in range(first) if any(LEADER.search(l) for l in pages[i].splitlines())]
    log("目录占 %d 页" % len(toc_idx))
    if len(toc_idx) < 2 or not shutil.which("pdftotext"):
        return
    try:
        out = subprocess.run(["pdftotext", "-bbox", "-f", str(toc_idx[0] + 1), "-l",
                              str(toc_idx[-1] + 1), str(pdf), "-"],
                             capture_output=True, text=True, timeout=60).stdout
    except Exception:
        return

    chunks = out.split("<page ")[1:]
    for k, chunk in enumerate(chunks[:-1]):
        height = float(re.search(r'height="([\d.]+)"', chunk).group(1))
        ys = [float(v) for v in re.findall(r'yMax="([\d.]+)"', chunk)]
        body = [y for y in ys if y < 0.9 * height]  # 排除页脚页码
        if body and max(body) / height < 0.75:
            log("警告：目录第 %d 页在 %.0f%% 处就换页，页尾留白过大，请检查是否有分页/分栏设置"
                % (k + 1, 100 * max(body) / height))


def fill_toc_pages(doc, entries, body_start_text, scheme, tmpdir):
    tmp_docx = Path(tmpdir) / "pass1.docx"
    doc.save(tmp_docx)
    pdf = render_pdf(tmp_docx, tmpdir)
    pages = pdf_pages(pdf) if pdf else None
    if not pages:
        log("警告：无法渲染 PDF，目录页码未回填，将设置为 Word 打开时更新")
        return False
    first = find_page(pages, body_start_text, 0)
    if first is None:
        multicol = body_is_multicol(doc)
        if multicol:
            log("正文为多栏排版，pdftotext 无法可靠定位标题；"
                "目录页码跳过回填，已设置 Word 打开时自动更新"
                "（在 Word 里右键目录 → 更新域 → 更新整个目录）")
        else:
            log("警告：找不到正文起始页，目录页码未回填；"
                "请确认正文首个 Heading 样式段落内容与原稿一致")
        return False
    check_toc_layout(pdf, pages, first)
    labels, cur = [], first
    for e in entries:
        idx = find_page(pages, e["text"], cur)
        if idx is None:
            log("警告：渲染结果中找不到标题\"%s\"，目录页码未回填" % e["text"])
            return False
        cur = idx
        labels.append(idx + 1 if scheme == "continuous" else idx - first + 1)
    for e, label in zip(entries, labels):
        e["t"].text = str(label)
    log("目录页码已按渲染结果回填（共 %d 条，PDF 共 %d 页）" % (len(entries), len(pages)))
    return True


# --------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description="中文课程报告 Word 排版")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--defaults", action="store_true")
    mode.add_argument("--black-only", action="store_true")
    parser.add_argument("--three-line", help="1-based comma-separated table indices, e.g. 1,3")
    parser.add_argument("--fix-abstract", action="store_true")
    parser.add_argument("--cover", help="cover.json 路径")
    parser.add_argument("--existing-cover", action="store_true")
    parser.add_argument("--keep-front", action="store_true")

    parser.add_argument("--toc", action="store_true")
    parser.add_argument("--toc-levels", type=int, default=3)
    parser.add_argument("--page-numbers", nargs="?", const="roman", choices=["roman", "continuous"])
    parser.add_argument("--no-fill-toc", action="store_true")
    parser.add_argument("--timezone", default="Asia/Taipei")
    opts = parser.parse_args()

    doc = Document(opts.source)
    gen, gen_refs = set(), []

    if opts.defaults:
        for section in doc.sections:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.top_margin = section.bottom_margin = Cm(2.5)
            section.left_margin = section.right_margin = Cm(2.5)
        fix_doc_defaults(doc)
        define_styles(doc, list(SPEC) + list(EXTRA))

    if opts.defaults:
        normalize_styles(doc)
        tidy_code_style(doc)

    if opts.three_line:  # 序号按原稿表格顺序，所以放在任何结构改动之前
        if opts.three_line.strip().lower() == "all":
            indices = set(range(1, len(doc.tables) + 1))
        else:
            indices = {int(item) for item in opts.three_line.split(",")}
        if indices and (min(indices) < 1 or max(indices) > len(doc.tables)):
            parser.error("--three-line indices must be existing 1-based table numbers")
        for index in indices:
            three_line_table(doc.tables[index - 1])


    fix_caption_positions(doc)
    tidy_tables(doc)

    entries, body_start_text = [], ""
    structural = opts.cover or opts.toc or opts.page_numbers or opts.existing_cover
    if opts.fix_abstract and not structural:
        fix_abstract(doc)
    if structural:
        define_styles(doc, ["Footer"])
        entries, body_start_text = build_structure(doc, opts, gen, gen_refs)

    names = set(ALL_SPEC) if opts.defaults else set(EXTRA)
    apply_all(doc, names, gen, opts.defaults)
    if opts.defaults:
        format_references(doc)
    tidy_figures(doc, gen)

    if entries:
        ok = False
        if not opts.no_fill_toc:
            if body_columns(doc) > 1:
                log("正文为多栏排版，目录页码未回填；已设置 Word 打开时自动更新"
                    "（在 Word 里右键目录 → 更新域 → 更新整个目录）")
            else:
                with tempfile.TemporaryDirectory() as tmp:
                    ok = fill_toc_pages(doc, entries, body_start_text,
                                        opts.page_numbers or "roman", tmp)
        if not ok:
            set_update_fields(doc)

    force_black(doc)
    opts.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(opts.output)
    log("完成：" + str(opts.output))


if __name__ == "__main__":
    main()
